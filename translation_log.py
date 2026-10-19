"""Append each translate run to a shared Excel workbook."""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from config import PROJECT_ROOT, TRANSLATION_LOG_PATH

logger = logging.getLogger(__name__)

_HEADERS = ("timestamp", "source_lang", "target_lang", "input_text", "translation")


class TranslationLogError(Exception):
    """Raised when the translation Excel log cannot be updated."""


def _resolve_log_path(log_path: Path | None) -> Path:
    path = Path(log_path) if log_path is not None else Path(TRANSLATION_LOG_PATH)
    if not path.is_absolute():
        path = PROJECT_ROOT / path
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def _ensure_workbook(path: Path):
    if path.exists():
        wb = load_workbook(path)
        ws = wb.active
        if ws.max_row == 0 or (ws.max_row == 1 and all(c.value is None for c in ws[1])):
            ws.append(list(_HEADERS))
        return wb, ws
    wb = Workbook()
    ws = wb.active
    ws.title = "translations"
    ws.append(list(_HEADERS))
    return wb, ws


def append_translation(
    input_text: str,
    translation: str,
    source_lang: str,
    target_lang: str,
    *,
    log_path: Path | None = None,
) -> Path:
    """Append one input/translation pair and return the Excel path."""
    path = _resolve_log_path(log_path)

    try:
        wb, ws = _ensure_workbook(path)
        ws.append(
            [
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
                (source_lang or "").strip().lower(),
                (target_lang or "").strip().lower(),
                (input_text or "").strip(),
                (translation or "").strip(),
            ]
        )
        wb.save(path)
    except OSError as exc:
        raise TranslationLogError(f"Could not write Excel log to {path}: {exc}") from exc
    except Exception as exc:  # noqa: BLE001
        raise TranslationLogError(f"Excel log update failed: {exc}") from exc

    logger.info("Appended translation to %s", path)
    return path.resolve()
