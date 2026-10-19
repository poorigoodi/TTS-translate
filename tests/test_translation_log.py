"""Tests for the append-only Excel translation log."""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from translation_log import append_translation  # noqa: E402


def test_append_creates_and_grows(tmp_path):
    path = tmp_path / "log.xlsx"
    append_translation("Hello", "سلام", "en", "fa", log_path=path)
    append_translation("Good morning", "صبح بخیر", "en", "fa", log_path=path)

    wb = load_workbook(path)
    ws = wb.active
    assert [c.value for c in ws[1]] == [
        "timestamp",
        "source_lang",
        "target_lang",
        "input_text",
        "translation",
    ]
    assert ws.max_row == 3
    assert ws.cell(2, 4).value == "Hello"
    assert ws.cell(2, 5).value == "سلام"
    assert ws.cell(3, 4).value == "Good morning"
    assert ws.cell(3, 5).value == "صبح بخیر"
